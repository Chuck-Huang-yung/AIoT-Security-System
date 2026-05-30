"""產生 AIoT 專題報告 PPT — 完整重製版"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ── 配色 ──
BG = RGBColor(0x0F, 0x17, 0x2A)
CARD = RGBColor(0x1E, 0x29, 0x3B)
CODE_BG = RGBColor(0x0D, 0x11, 0x17)
CODE_BORDER = RGBColor(0x30, 0x3B, 0x50)
BLUE = RGBColor(0x38, 0xBD, 0xF8)
GREEN = RGBColor(0x4A, 0xDE, 0x80)
RED = RGBColor(0xF8, 0x71, 0x71)
YELLOW = RGBColor(0xFB, 0xBF, 0x24)
PURPLE = RGBColor(0xA7, 0x8B, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x94, 0xA3, 0xB8)
LGRAY = RGBColor(0xCB, 0xD5, 0xE1)

FONT = "Microsoft JhengHei"
MONO = "Consolas"

# ── 簡報初始化 ──
prs = Presentation()
W = Inches(13.333)
H = Inches(7.5)
prs.slide_width = W
prs.slide_height = H


# ── 工具函式 ──
def new_slide():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    return s


def rect(slide, l, t, w, h, fill=CARD, border=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if border:
        sh.line.color.rgb = border
        sh.line.width = Pt(1.5)
    else:
        sh.line.fill.background()
    return sh


def txt(slide, l, t, w, h, text, sz=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT, font=FONT):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(sz)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font
    p.alignment = align
    return tb


def bullets(slide, l, t, w, h, items, sz=16, color=WHITE, dot_color=BLUE, spacing=10):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(spacing)
        r1 = p.add_run()
        r1.text = "●  "
        r1.font.size = Pt(sz - 2)
        r1.font.color.rgb = dot_color
        r1.font.name = FONT
        r2 = p.add_run()
        r2.text = item
        r2.font.size = Pt(sz)
        r2.font.color.rgb = color
        r2.font.name = FONT
    return tb


def code(slide, l, t, w, h, text, sz=11):
    sh = rect(slide, l, t, w, h, fill=CODE_BG, border=CODE_BORDER)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(14)
    tf.margin_right = Pt(14)
    tf.margin_top = Pt(10)
    tf.margin_bottom = Pt(10)
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        r = p.add_run()
        r.text = line
        r.font.size = Pt(sz)
        r.font.color.rgb = GREEN
        r.font.name = MONO
    return sh


def title_bar(slide, title, sub=None):
    txt(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
        title, sz=30, bold=True)
    if sub:
        txt(slide, Inches(0.8), Inches(0.95), Inches(11), Inches(0.4),
            sub, sz=15, color=GRAY)
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                Inches(0.8), Inches(1.4), Inches(1.8), Pt(3))
    ln.fill.solid()
    ln.fill.fore_color.rgb = BLUE
    ln.line.fill.background()


# ── 流程步驟列（水平卡片） ──
def step_row(slide, steps, y_top=Inches(2.0), card_w=Inches(1.75), card_h=Inches(1.9), gap=Inches(0.25)):
    n = len(steps)
    total = card_w * n + gap * (n - 1)
    x0 = (W - total) / 2  # 置中
    for i, (num, title, desc, color) in enumerate(steps):
        x = int(x0 + i * (card_w + gap))
        rect(slide, x, y_top, card_w, card_h, border=color)
        # 數字
        csz = Inches(0.45)
        cx = x + (card_w - csz) // 2
        circle = rect(slide, cx, y_top + Inches(0.12), csz, csz, fill=color)
        tf = circle.text_frame
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tf.paragraphs[0].add_run()
        r.text = num
        r.font.size = Pt(18)
        r.font.color.rgb = BG
        r.font.bold = True
        r.font.name = MONO
        # 標題
        txt(slide, x + Inches(0.05), y_top + Inches(0.65), card_w - Inches(0.1), Inches(0.3),
            title, sz=13, color=color, bold=True, align=PP_ALIGN.CENTER)
        # 描述
        txt(slide, x + Inches(0.05), y_top + Inches(1.0), card_w - Inches(0.1), Inches(0.8),
            desc, sz=11, color=LGRAY, align=PP_ALIGN.CENTER)
        # 箭頭
        if i < n - 1:
            txt(slide, x + card_w, y_top + Inches(0.7), gap, Inches(0.3),
                "→", sz=18, color=GRAY, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════
# Slide 1 — 封面
# ══════════════════════════════════════════════════════════════
s = new_slide()
txt(s, Inches(0.5), Inches(1.8), Inches(12.3), Inches(1),
    "AIoT 智慧安防與環境感測系統",
    sz=44, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(0.5), Inches(3.0), Inches(12.3), Inches(0.6),
    "基於 Raspberry Pi + ESP32 的人體偵測、即時監控與自動照明控制平台",
    sz=20, color=GRAY, align=PP_ALIGN.CENTER)

tags = [("Raspberry Pi", BLUE), ("ESP32", GREEN), ("MQTT", PURPLE)]
for i, (t, c) in enumerate(tags):
    x = Inches(4.0 + i * 2.0)
    sh = rect(s, x, Inches(4.0), Inches(1.7), Inches(0.5), border=c)
    tf = sh.text_frame
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].add_run()
    r.text = t
    r.font.size = Pt(14)
    r.font.color.rgb = c
    r.font.bold = True
    r.font.name = MONO

txt(s, Inches(0.5), Inches(5.8), Inches(12.3), Inches(0.4),
    "AIoT 期末專題", sz=16, color=GRAY, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════
# Slide 2 — 動機與目的
# ══════════════════════════════════════════════════════════════
s = new_slide()
title_bar(s, "專題動機與目的")

bullets(s, Inches(0.8), Inches(1.8), Inches(5.5), Inches(3), [
    "IoT 技術普及，智慧家庭安防成為重要應用",
    "結合 AI 影像辨識與 IoT 感測技術",
    "Raspberry Pi + ESP32 雙裝置分散式架構",
    "雲端 MQTT Broker 實現跨裝置通訊",
], sz=17)

feats = [
    ("入侵偵測與警報", "PIR → 攝影機 → MobileNet SSD\n→ Discord 即時通知", RED),
    ("智慧自動照明", "PIR + 光敏電阻聯動判斷\n→ 有人且暗時自動開燈", YELLOW),
    ("Web 即時監控", "MQTT.js 直連 Broker\n→ 即時數據 + 遠端控制", BLUE),
]
for i, (ft, desc, c) in enumerate(feats):
    y = Inches(1.8 + i * 1.7)
    rect(s, Inches(7.0), y, Inches(5.5), Inches(1.4), border=c)
    txt(s, Inches(7.4), y + Inches(0.15), Inches(4.8), Inches(0.35),
        ft, sz=17, color=c, bold=True)
    txt(s, Inches(7.4), y + Inches(0.55), Inches(4.8), Inches(0.7),
        desc, sz=13, color=LGRAY)


# ══════════════════════════════════════════════════════════════
# Slide 3 — 系統架構
# ══════════════════════════════════════════════════════════════
s = new_slide()
title_bar(s, "系統架構總覽")

# ESP32
rect(s, Inches(0.5), Inches(1.7), Inches(3.3), Inches(5.0), border=GREEN)
txt(s, Inches(0.8), Inches(1.8), Inches(2.8), Inches(0.35),
    "ESP32（MicroPython）", sz=15, color=GREEN, bold=True)
txt(s, Inches(0.8), Inches(2.15), Inches(2.8), Inches(0.25),
    "感測 + 照明控制", sz=12, color=GRAY)
bullets(s, Inches(0.8), Inches(2.5), Inches(2.8), Inches(3.5), [
    "PIR 感測器 (GPIO 27)",
    "光敏電阻 ADC (GPIO 34)",
    "LED 自動/手動控制",
    "MQTT 發布 PIR + 亮度",
    "MQTT 訂閱控燈指令",
], sz=12, dot_color=GREEN, spacing=6)

# RPi
rect(s, Inches(0.5), Inches(4.6), Inches(3.3), Inches(2.1), border=BLUE)
txt(s, Inches(0.8), Inches(4.7), Inches(2.8), Inches(0.35),
    "Raspberry Pi（Python）", sz=15, color=BLUE, bold=True)
txt(s, Inches(0.8), Inches(5.05), Inches(2.8), Inches(0.25),
    "安防偵測線路", sz=12, color=GRAY)
bullets(s, Inches(0.8), Inches(5.35), Inches(2.8), Inches(1.5), [
    "USB Webcam + OpenCV DNN",
    "MobileNet SSD 人形偵測",
    "Discord 警報推播",
], sz=12, dot_color=BLUE, spacing=5)

# MQTT Broker
rect(s, Inches(4.8), Inches(2.8), Inches(3.5), Inches(2.2), border=PURPLE)
txt(s, Inches(5.1), Inches(2.9), Inches(3.0), Inches(0.35),
    "HiveMQ Cloud", sz=17, color=PURPLE, bold=True)
txt(s, Inches(5.1), Inches(3.3), Inches(3.0), Inches(0.25),
    "MQTT Broker（雲端）", sz=12, color=GRAY)
bullets(s, Inches(5.1), Inches(3.65), Inches(3.0), Inches(1.2), [
    "TLS 加密 (port 8883)",
    "WSS 連線 (port 8884)",
    "免費方案，外網可存取",
], sz=11, dot_color=PURPLE, spacing=5)

# Dashboard
rect(s, Inches(9.3), Inches(1.7), Inches(3.5), Inches(2.5), border=YELLOW)
txt(s, Inches(9.6), Inches(1.8), Inches(3.0), Inches(0.35),
    "Web Dashboard", sz=15, color=YELLOW, bold=True)
txt(s, Inches(9.6), Inches(2.15), Inches(3.0), Inches(0.25),
    "即時監控介面", sz=12, color=GRAY)
bullets(s, Inches(9.6), Inches(2.5), Inches(3.0), Inches(1.5), [
    "即時截圖 + PIR 狀態",
    "亮度數值 + 燈光控制",
    "警報事件 + 系統狀態",
], sz=12, dot_color=YELLOW, spacing=5)

# Discord
rect(s, Inches(9.3), Inches(4.5), Inches(3.5), Inches(1.3), border=RED)
txt(s, Inches(9.6), Inches(4.6), Inches(3.0), Inches(0.35),
    "Discord 通知", sz=15, color=RED, bold=True)
txt(s, Inches(9.6), Inches(5.0), Inches(3.0), Inches(0.6),
    "Webhook 即時推播\n入侵警報 + 截圖", sz=12, color=LGRAY)

# 箭頭
for x, y, t, c in [
    (Inches(3.9), Inches(3.2), "MQTT →", GREEN),
    (Inches(3.9), Inches(5.2), "MQTT →", BLUE),
    (Inches(8.4), Inches(2.5), "→ WSS", YELLOW),
    (Inches(8.4), Inches(4.8), "→ HTTP", RED),
]:
    txt(s, x, y, Inches(0.9), Inches(0.3), t, sz=11, color=c, bold=True)


# ══════════════════════════════════════════════════════════════
# Slide 4 — 安防偵測流程
# ══════════════════════════════════════════════════════════════
s = new_slide()
title_bar(s, "線路一：安防偵測流程",
          "ESP32 PIR → MQTT → RPi Camera → MobileNet SSD → Discord")

step_row(s, [
    ("1", "ESP32 PIR", "紅外線感測器\nGPIO 27 偵測", GREEN),
    ("2", "MQTT 觸發", "發布 PIR 訊號\nRPi 接收觸發", PURPLE),
    ("3", "連拍驗證", "USB Webcam\n連拍 3 張", BLUE),
    ("4", "MobileNet SSD", "深度學習模型\n2/3 通過才算", YELLOW),
    ("5", "MQTT 發布", "警報事件\n+ 截圖影像", BLUE),
    ("6", "Discord 通知", "Webhook 推播\n含入侵截圖", RED),
], y_top=Inches(1.8))

rect(s, Inches(0.8), Inches(4.2), Inches(11.7), Inches(1.0), border=YELLOW)
txt(s, Inches(1.1), Inches(4.3), Inches(11.0), Inches(0.3),
    "防抖動機制", sz=15, color=YELLOW, bold=True)
txt(s, Inches(1.1), Inches(4.65), Inches(11.0), Inches(0.5),
    "多幀驗證：連拍 3 張，至少 2 張偵測到人形才觸發警報。RPi 冷卻 15 秒，ESP32 PIR 冷卻 10 秒。",
    sz=13, color=LGRAY)


# ══════════════════════════════════════════════════════════════
# Slide 5 — 智慧照明流程
# ══════════════════════════════════════════════════════════════
s = new_slide()
title_bar(s, "線路二：智慧照明流程",
          "ESP32 — PIR + 光敏電阻 → 智慧 LED 控制 → MQTT")

step_row(s, [
    ("1", "PIR 偵測", "GPIO 27\n是否有人在場", BLUE),
    ("2", "ADC 讀取", "光敏電阻\n亮度 0-4095", GREEN),
    ("3", "智慧判斷", "有人+暗→開燈\n無人→關燈", YELLOW),
    ("4", "MQTT 發布", "亮度 + 燈光\n+ PIR 狀態", PURPLE),
    ("5", "遠端控制", "Dashboard 手動\n開/關燈", RED),
    ("6", "手動覆蓋", "手動控制後\n暫停自動 60s", BLUE),
], y_top=Inches(1.8))

rect(s, Inches(0.8), Inches(4.2), Inches(11.7), Inches(1.0), border=GREEN)
txt(s, Inches(1.1), Inches(4.3), Inches(11.0), Inches(0.3),
    "智慧照明邏輯", sz=15, color=GREEN, bold=True)
txt(s, Inches(1.1), Inches(4.65), Inches(11.0), Inches(0.5),
    "有人時依亮度決定是否開燈（供攝影機夜間辨識），無人時自動關燈。Dashboard 手動控燈後暫停自動模式 60 秒。",
    sz=13, color=LGRAY)


# ══════════════════════════════════════════════════════════════
# Slide 6 — MQTT 通訊架構
# ══════════════════════════════════════════════════════════════
s = new_slide()
title_bar(s, "MQTT 通訊架構", "Topic 設計與訊息格式")

topics = [
    ("home/security/pir",      "ESP32",     "RPi, Dashboard", '{"triggered":true}',        "PIR 觸發",  BLUE),
    ("home/security/alert",    "RPi",       "Dashboard",      '{"confidence":0.85}',       "入侵警報",  RED),
    ("home/security/snapshot", "RPi",       "Dashboard",      '{"image":"<Base64>"}',      "截圖影像",  RED),
    ("home/sensor/light",      "ESP32",     "Dashboard",      '{"value":2048}',            "亮度數值",  GREEN),
    ("home/light/status",      "ESP32",     "Dashboard",      '{"status":"on"}',           "燈光狀態",  GREEN),
    ("home/light/control",     "Dashboard", "ESP32",          '{"action":"on"}',           "遠端控燈",  YELLOW),
    ("home/system/status",     "RPi",       "Dashboard",      '{"cpu_temp":45}',           "系統狀態",  BLUE),
]

# 表頭
hdrs = ["Topic", "發布者", "訂閱者", "Payload 範例", "說明"]
ws = [Inches(2.6), Inches(1.2), Inches(1.6), Inches(2.8), Inches(1.5)]
x0 = Inches(0.8)
y0 = Inches(1.7)

xp = x0
for h, w in zip(hdrs, ws):
    txt(s, xp, y0, w, Inches(0.3), h, sz=13, color=BLUE, bold=True)
    xp += w

# 分隔線
ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x0, Inches(2.0), sum(ws), Pt(1))
ln.fill.solid()
ln.fill.fore_color.rgb = CODE_BORDER
ln.line.fill.background()

# 內容
for ri, (topic, pub, sub, payload, desc, c) in enumerate(topics):
    y = Inches(2.15 + ri * 0.7)
    row = [topic, pub, sub, payload, desc]
    xp = x0
    for ci, (d, w) in enumerate(zip(row, ws)):
        clr = GREEN if ci == 0 else (c if ci == 4 else LGRAY)
        fsz = 10 if ci in (0, 3) else 12
        fn = MONO if ci in (0, 3) else FONT
        txt(s, xp, y, w, Inches(0.6), d, sz=fsz, color=clr, font=fn)
        xp += w


# ══════════════════════════════════════════════════════════════
# Slide 7 — 硬體配置
# ══════════════════════════════════════════════════════════════
s = new_slide()
title_bar(s, "硬體配置")

# RPi
rect(s, Inches(0.5), Inches(1.7), Inches(5.8), Inches(5.0), border=BLUE)
txt(s, Inches(0.8), Inches(1.8), Inches(5.0), Inches(0.4),
    "Raspberry Pi 端", sz=20, color=BLUE, bold=True)
bullets(s, Inches(0.8), Inches(2.4), Inches(5.2), Inches(2.0), [
    "Raspberry Pi 4 Model B — 安防主控",
    "USB Webcam (Logitech C270) — 影像擷取",
    "OpenCV DNN + MobileNet SSD 人形偵測",
    "供電：5V USB-C",
], sz=14, dot_color=BLUE, spacing=8)
code(s, Inches(0.8), Inches(4.5), Inches(5.2), Inches(1.8),
     "# 無需 GPIO 感測器\n# PIR 觸發訊號由 ESP32\n# 透過 MQTT 遠端傳送\n\nmqtt.subscribe('home/security/pir',\n               on_pir_trigger)", sz=11)

# ESP32
rect(s, Inches(7.0), Inches(1.7), Inches(5.8), Inches(5.0), border=GREEN)
txt(s, Inches(7.3), Inches(1.8), Inches(5.0), Inches(0.4),
    "ESP32 端", sz=20, color=GREEN, bold=True)
bullets(s, Inches(7.3), Inches(2.4), Inches(5.2), Inches(2.0), [
    "ESP32-WROOM-32 — 感測 + 照明主控",
    "PIR 紅外線感測器 (SR505) — GPIO 27",
    "光敏電阻 (LDR) — GPIO 34 (ADC)",
    "LED — GPIO 2",
], sz=14, dot_color=GREEN, spacing=8)
code(s, Inches(7.3), Inches(4.5), Inches(5.2), Inches(1.8),
     "# PIR 接線\nOUT → GPIO 27,  VCC → 5V,  GND → GND\n\n# LDR 接線\n一端 → 3.3V,  另一端 → GPIO 34 + 10K\n\n# LED 接線\n長腳 → GPIO 2（經電阻）,  短腳 → GND", sz=10)


# ══════════════════════════════════════════════════════════════
# Slide 8 — 軟體技術棧
# ══════════════════════════════════════════════════════════════
s = new_slide()
title_bar(s, "軟體技術棧")

cats = [
    ("Raspberry Pi", BLUE, [
        ("Python 3", "主程式語言"),
        ("OpenCV DNN", "MobileNet SSD 人形偵測"),
        ("paho-mqtt v2", "MQTT Client（TLS 加密）"),
        ("requests", "Discord Webhook 請求"),
        ("threading", "多線程（安防 + 系統狀態）"),
    ]),
    ("ESP32", GREEN, [
        ("MicroPython", "統一 Python 生態"),
        ("umqtt.simple", "輕量 MQTT Client（SSL）"),
        ("machine.ADC", "光敏電阻類比讀取"),
        ("machine.Pin", "GPIO 控制（PIR, LED）"),
    ]),
    ("雲端 / 前端", PURPLE, [
        ("HiveMQ Cloud", "MQTT Broker（免費雲端）"),
        ("MQTT.js", "前端 WSS 直連 Broker"),
        ("Tailwind CSS", "現代深色主題 UI"),
        ("Discord Webhook", "入侵警報推播通知"),
    ]),
]

col_w = Inches(3.8)
col_gap = Inches(0.3)
x_start = Inches(0.5)

for ci, (name, c, techs) in enumerate(cats):
    x = x_start + ci * (col_w + col_gap)
    rect(s, x, Inches(1.7), col_w, Inches(5.2), border=c)
    txt(s, x + Inches(0.3), Inches(1.85), col_w - Inches(0.6), Inches(0.35),
        name, sz=17, color=c, bold=True)
    for ti, (tech, desc) in enumerate(techs):
        y = Inches(2.4 + ti * 0.9)
        txt(s, x + Inches(0.3), y, col_w - Inches(0.6), Inches(0.3),
            tech, sz=14, color=WHITE, bold=True)
        txt(s, x + Inches(0.3), y + Inches(0.3), col_w - Inches(0.6), Inches(0.3),
            desc, sz=11, color=GRAY)


# ══════════════════════════════════════════════════════════════
# Slide 9 — RPi 安防主迴圈
# ══════════════════════════════════════════════════════════════
s = new_slide()
title_bar(s, "核心程式碼 — RPi 安防主迴圈",
          "main.py — MQTT 觸發 + 多幀驗證")

code(s, Inches(0.5), Inches(1.6), Inches(7.8), Inches(5.4),
     'pir_triggered = threading.Event()\n'
     '\n'
     'def on_pir_trigger(topic, payload):\n'
     '    pir_triggered.set()  # ESP32 PIR 觸發\n'
     '\n'
     'def security_loop():\n'
     '    capture.init()\n'
     '    COOLDOWN = 15\n'
     '\n'
     '    while running:\n'
     '        if pir_triggered.wait(timeout=1):\n'
     '            pir_triggered.clear()\n'
     '\n'
     '            # 多幀驗證：連拍 3 張\n'
     '            detect_count = 0\n'
     '            for attempt in range(3):\n'
     '                frame = capture.capture_frame()\n'
     '                detected, boxes, conf = \\\n'
     '                    detector.detect_person(frame)\n'
     '                if detected:\n'
     '                    detect_count += 1\n'
     '                time.sleep(0.3)\n'
     '\n'
     '            if detect_count >= 2:  # 2/3 通過\n'
     '                mqtt_client.publish(ALERT, {...})\n'
     '                discord_bot.send_alert(\n'
     '                    f"偵測到入侵！{conf:.0%}",\n'
     '                    image_base64=img_base64)\n'
     '            time.sleep(COOLDOWN)', sz=11)

notes = [
    ("MQTT 觸發", "ESP32 PIR 偵測移動\n透過 MQTT 遠端觸發 RPi", GREEN),
    ("多幀驗證", "連拍 3 張照片\n至少 2/3 通過才觸發警報", YELLOW),
    ("MobileNet SSD", "Caffe 深度學習模型\n精準人形偵測 (class 15)", PURPLE),
    ("多管道通知", "MQTT → Dashboard 即時\nWebhook → Discord 推播", RED),
]
for i, (t, d, c) in enumerate(notes):
    y = Inches(1.7 + i * 1.3)
    rect(s, Inches(8.8), y, Inches(4.0), Inches(1.1), border=c)
    txt(s, Inches(9.1), y + Inches(0.1), Inches(3.5), Inches(0.3),
        t, sz=13, color=c, bold=True)
    txt(s, Inches(9.1), y + Inches(0.4), Inches(3.5), Inches(0.6),
        d, sz=11, color=LGRAY)


# ══════════════════════════════════════════════════════════════
# Slide 10 — ESP32 PIR + 智慧照明
# ══════════════════════════════════════════════════════════════
s = new_slide()
title_bar(s, "核心程式碼 — ESP32 PIR + 智慧照明",
          "esp32/main.py — MicroPython")

code(s, Inches(0.5), Inches(1.6), Inches(7.8), Inches(5.4),
     '# 硬體初始化\n'
     'pir = Pin(27, Pin.IN, Pin.PULL_DOWN)\n'
     'ldr = ADC(Pin(34))\n'
     'ldr.atten(ADC.ATTN_11DB)  # 0-3.3V\n'
     'led = Pin(2, Pin.OUT)\n'
     '\n'
     'while True:\n'
     '    ldr_value = ldr.read()\n'
     '    pir_value = pir.value()\n'
     '    now = time.time()\n'
     '\n'
     '    # 手動覆蓋超時 → 恢復自動\n'
     '    if manual_override and \\\n'
     '       (now - override_time > 60):\n'
     '        manual_override = False\n'
     '\n'
     '    # 智慧照明（手動模式下跳過）\n'
     '    if not manual_override:\n'
     '        if pir_value == 1:\n'
     '            if ldr_value < THRESHOLD:\n'
     '                set_led(client, True)\n'
     '                # 有人+暗 → 開燈\n'
     '        else:\n'
     '            set_led(client, False)\n'
     '            # 沒人 → 關燈\n'
     '\n'
     '    # PIR → MQTT 通知 RPi 拍照\n'
     '    if pir_value == 1 and \\\n'
     '       (now - last_pir >= COOLDOWN):\n'
     '        client.publish(\n'
     '            "home/security/pir",\n'
     '            \'{"triggered":true}\')', sz=11)

notes2 = [
    ("PIR 偵測", "GPIO 27 偵測人體移動\n觸發 MQTT 通知 RPi 拍照", GREEN),
    ("智慧照明", "有人 + 暗 → 開燈\n無人 → 關燈（省電）", YELLOW),
    ("手動覆蓋", "Dashboard 控制後\n暫停自動模式 60 秒", PURPLE),
    ("雙重角色", "感測端 + 照明控制\n單一 ESP32 完成", BLUE),
]
for i, (t, d, c) in enumerate(notes2):
    y = Inches(1.7 + i * 1.3)
    rect(s, Inches(8.8), y, Inches(4.0), Inches(1.1), border=c)
    txt(s, Inches(9.1), y + Inches(0.1), Inches(3.5), Inches(0.3),
        t, sz=13, color=c, bold=True)
    txt(s, Inches(9.1), y + Inches(0.4), Inches(3.5), Inches(0.6),
        d, sz=11, color=LGRAY)


# ══════════════════════════════════════════════════════════════
# Slide 11 — Web Dashboard
# ══════════════════════════════════════════════════════════════
s = new_slide()
title_bar(s, "Web Dashboard 即時監控",
          "Tailwind CSS + MQTT.js — 純靜態，無需後端")

feats = [
    ("即時截圖",     "subscribe snapshot\nBase64 圖片即時渲染", RED),
    ("PIR 感測器",   "subscribe pir\n即時顯示 PIR 狀態 0/1", BLUE),
    ("警報事件",     "subscribe alert\n入侵警報列表顯示", RED),
    ("環境亮度",     "subscribe light\nESP32 ADC 即時數據", GREEN),
    ("燈光控制",     "publish control\n手動開/關燈按鈕", YELLOW),
    ("系統狀態",     "subscribe status\nCPU 溫度、運行時間", PURPLE),
]

for i, (t, d, c) in enumerate(feats):
    col = i % 3
    row = i // 3
    x = Inches(0.5 + col * 4.15)
    y = Inches(1.7 + row * 2.3)
    rect(s, x, y, Inches(3.85), Inches(1.9), border=c)
    txt(s, x + Inches(0.3), y + Inches(0.2), Inches(3.3), Inches(0.35),
        t, sz=17, color=c, bold=True)
    txt(s, x + Inches(0.3), y + Inches(0.65), Inches(3.3), Inches(0.9),
        d, sz=13, color=LGRAY)

rect(s, Inches(0.5), Inches(6.3), Inches(12.3), Inches(0.8), border=BLUE)
txt(s, Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.6),
    "MQTT.js 透過 WSS 直連 HiveMQ Cloud，無需後端。純靜態 HTML/JS，可部署於任何 Web Server 或直接本地開啟。",
    sz=13, color=LGRAY)


# ══════════════════════════════════════════════════════════════
# Slide 12 — Discord 通知
# ══════════════════════════════════════════════════════════════
s = new_slide()
title_bar(s, "Discord 即時通知",
          "Webhook — 入侵警報 + 截圖推播")

code(s, Inches(0.5), Inches(1.6), Inches(6.8), Inches(4.5),
     'def send_alert(message, image_base64=None):\n'
     '    embed = {\n'
     '        "title": "入侵警報",\n'
     '        "description": message,\n'
     '        "color": 0xFF0000,\n'
     '        "timestamp": datetime.now().isoformat(),\n'
     '        "image": {\n'
     '            "url": "attachment://snapshot.jpg"\n'
     '        },\n'
     '    }\n'
     '\n'
     '    # multipart/form-data 上傳圖片\n'
     '    image_data = base64.b64decode(image_base64)\n'
     '    files = {\n'
     '        "payload_json": (\n'
     '            None, json.dumps(payload)),\n'
     '        "files[0]": (\n'
     '            "snapshot.jpg", BytesIO(image_data)),\n'
     '    }\n'
     '    requests.post(WEBHOOK_URL, files=files)', sz=11)

rect(s, Inches(7.8), Inches(1.6), Inches(5.0), Inches(4.5), border=RED)
txt(s, Inches(8.1), Inches(1.7), Inches(4.4), Inches(0.4),
    "通知功能", sz=18, color=RED, bold=True)
bullets(s, Inches(8.1), Inches(2.2), Inches(4.4), Inches(3.5), [
    "Embed 格式：標題 + 描述 + 時間戳",
    "附帶入侵截圖（JPG 圖片）",
    "multipart/form-data 上傳",
    "系統上線 / 下線通知",
    "無需 Bot Token，僅用 Webhook URL",
], sz=14, dot_color=RED, spacing=10)


# ══════════════════════════════════════════════════════════════
# Slide 13 — 專案目錄結構
# ══════════════════════════════════════════════════════════════
s = new_slide()
title_bar(s, "專案目錄結構")

code(s, Inches(0.5), Inches(1.6), Inches(7.3), Inches(5.5),
     'AIOT_final/\n'
     '├── config.py              # RPi 設定檔\n'
     '├── main.py                # RPi 主程式入口\n'
     '├── .env                   # 敏感資訊\n'
     '│\n'
     '├── camera/\n'
     '│   ├── capture.py         # 攝影機擷取\n'
     '│   ├── detector.py        # MobileNet SSD\n'
     '│   └── model/\n'
     '│       ├── deploy.prototxt\n'
     '│       └── mobilenet_ssd.caffemodel\n'
     '│\n'
     '├── mqtt/\n'
     '│   └── client.py          # paho-mqtt v2\n'
     '│\n'
     '├── notify/\n'
     '│   └── discord_bot.py     # Discord Webhook\n'
     '│\n'
     '├── web/\n'
     '│   └── index.html         # Dashboard\n'
     '│\n'
     '├── esp32/                  # MicroPython\n'
     '│   ├── boot.py            # WiFi 連線\n'
     '│   ├── main.py            # PIR+LDR+LED\n'
     '│   ├── config.py\n'
     '│   └── lib/umqtt_simple.py\n'
     '│\n'
     '└── snapshots/', sz=11)

mods = [
    ("RPi 端", "安防偵測 + 系統監控\nPython 3 + OpenCV DNN\n+ paho-mqtt v2", BLUE),
    ("ESP32 端", "PIR 感測 + 照明控制\nMicroPython + umqtt", GREEN),
    ("前端", "即時監控 Dashboard\nTailwind + MQTT.js", YELLOW),
    ("雲端", "HiveMQ Cloud Broker\n+ Discord Webhook", PURPLE),
]
for i, (t, d, c) in enumerate(mods):
    y = Inches(1.7 + i * 1.35)
    rect(s, Inches(8.3), y, Inches(4.5), Inches(1.15), border=c)
    txt(s, Inches(8.6), y + Inches(0.1), Inches(4.0), Inches(0.3),
        t, sz=14, color=c, bold=True)
    txt(s, Inches(8.6), y + Inches(0.4), Inches(4.0), Inches(0.7),
        d, sz=12, color=LGRAY)


# ══════════════════════════════════════════════════════════════
# Slide 14 — 問題解決
# ══════════════════════════════════════════════════════════════
s = new_slide()
title_bar(s, "開發過程與問題解決")

probs = [
    ("PIR 持續誤觸發",
     "GPIO 浮接 → 持續高電位",
     "pull-down 電阻 + 軟體冷卻",
     RED),
    ("Haar Cascade 誤判衣服為人",
     "upperbody 偵測器誤報率高",
     "改用 MobileNet SSD 深度學習模型",
     YELLOW),
    ("PIR 與攝影機位置衝突",
     "同一位置角度難以兼顧",
     "PIR 移至 ESP32，MQTT 遠端觸發",
     GREEN),
    ("手動控燈被自動覆蓋",
     "自動模式立即覆蓋手動操作",
     "manual_override，暫停自動 60 秒",
     BLUE),
    ("ESP32 ussl 模組更名",
     "MicroPython v1.25 ussl → ssl",
     "umqtt_simple 加入 try/except 相容",
     PURPLE),
]

for i, (prob, cause, sol, c) in enumerate(probs):
    y = Inches(1.65 + i * 1.05)
    rect(s, Inches(0.5), y, Inches(12.3), Inches(0.9), border=c)
    txt(s, Inches(0.8), y + Inches(0.1), Inches(3.0), Inches(0.7),
        prob, sz=13, color=c, bold=True)
    txt(s, Inches(3.8), y + Inches(0.1), Inches(3.8), Inches(0.7),
        "原因：" + cause, sz=11, color=GRAY)
    txt(s, Inches(7.8), y + Inches(0.1), Inches(4.8), Inches(0.7),
        "解法：" + sol, sz=11, color=GREEN)


# ══════════════════════════════════════════════════════════════
# Slide 15 — 系統成果
# ══════════════════════════════════════════════════════════════
s = new_slide()
title_bar(s, "系統成果")

results = [
    ("安防偵測完整流程",
     "ESP32 PIR → MQTT → RPi 連拍 3 張\n→ MobileNet SSD 多幀驗證（2/3 通過）\n→ Discord 推播通知（含截圖）",
     RED),
    ("智慧照明控制",
     "PIR 偵測有人 + 光敏判斷環境暗\n→ 自動開燈（供攝影機夜間辨識）\n→ Dashboard 手動控制（60 秒覆蓋）",
     GREEN),
    ("Web Dashboard 即時監控",
     "MQTT.js 直連雲端 Broker\n→ 即時截圖、PIR 狀態、亮度、燈光控制\n→ 警報記錄、系統狀態監控",
     YELLOW),
    ("雲端分散式架構",
     "PIR 在 ESP32、攝影機在 RPi\n→ MQTT 解耦，支援外網存取\n→ systemd 開機自啟動",
     BLUE),
]

for i, (t, d, c) in enumerate(results):
    col = i % 2
    row = i // 2
    x = Inches(0.5 + col * 6.3)
    y = Inches(1.7 + row * 2.6)
    rect(s, x, y, Inches(5.9), Inches(2.3), border=c)
    txt(s, x + Inches(0.3), y + Inches(0.2), Inches(5.3), Inches(0.35),
        t, sz=17, color=c, bold=True)
    txt(s, x + Inches(0.3), y + Inches(0.65), Inches(5.3), Inches(1.4),
        d, sz=14, color=LGRAY)


# ══════════════════════════════════════════════════════════════
# Slide 16 — 結尾
# ══════════════════════════════════════════════════════════════
s = new_slide()
txt(s, Inches(0.5), Inches(2.2), Inches(12.3), Inches(1),
    "Thank You", sz=52, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(0.5), Inches(3.5), Inches(12.3), Inches(0.6),
    "AIoT 智慧安防與環境感測系統", sz=24, color=GRAY, align=PP_ALIGN.CENTER)

footer_tags = [
    ("Raspberry Pi", BLUE), ("ESP32", GREEN), ("OpenCV", PURPLE),
    ("MQTT", YELLOW), ("Discord", RED), ("Tailwind CSS", BLUE),
]
for i, (t, c) in enumerate(footer_tags):
    x = Inches(2.1 + i * 1.55)
    sh = rect(s, x, Inches(4.8), Inches(1.35), Inches(0.4), border=c)
    tf = sh.text_frame
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].add_run()
    r.text = t
    r.font.size = Pt(11)
    r.font.color.rgb = c
    r.font.bold = True
    r.font.name = MONO


# ══════════════════════════════════════════════════════════════
# 儲存
# ══════════════════════════════════════════════════════════════
out = "/Users/ianho/AIOT_final/AIoT_期末專題報告.pptx"
prs.save(out)
print(f"PPT 已儲存: {out}")
